#!/usr/bin/env python3
"""
=========================================================
  EXTRACTOR COMPLETO - productosdelimpiezalima.com
  WooCommerce REST API + HTML scraping
  Genera: products.xlsx  database.sql  images.json  sitemap.json
=========================================================
  USO:
    pip install -r requirements.txt
    python woo_extractor.py
=========================================================
"""

import os
import sys
import json
import time
import re
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from bs4 import BeautifulSoup
from tqdm import tqdm
from colorama import Fore, Style, init
from datetime import datetime

# ── config ────────────────────────────────────────────
sys.path.insert(0, os.path.dirname(__file__))
from config import (
    SITE_URL, WC_CONSUMER_KEY, WC_CONSUMER_SECRET,
    PER_PAGE, OUTPUT_DIR, REQUESTS_TIMEOUT, REQUESTS_RETRIES, REQUESTS_DELAY
)

init(autoreset=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

TIMESTAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

# ── helpers ───────────────────────────────────────────

def log(msg, color=Fore.CYAN):
    print(f"{color}[✓] {msg}{Style.RESET_ALL}")

def warn(msg):
    print(f"{Fore.YELLOW}[!] {msg}{Style.RESET_ALL}")

def error(msg):
    print(f"{Fore.RED}[✗] {msg}{Style.RESET_ALL}")

def wc_get(endpoint, params=None):
    """GET paginado a la WooCommerce REST API v3."""
    auth = (WC_CONSUMER_KEY, WC_CONSUMER_SECRET)
    url  = f"{SITE_URL}/wp-json/wc/v3/{endpoint}"
    all_items = []
    page = 1

    bar = tqdm(desc=f"  {endpoint}", unit="items", leave=False)
    while True:
        p = {"per_page": PER_PAGE, "page": page, **(params or {})}
        for attempt in range(REQUESTS_RETRIES):
            try:
                r = requests.get(url, auth=auth, params=p, timeout=REQUESTS_TIMEOUT)
                r.raise_for_status()
                break
            except Exception as e:
                if attempt == REQUESTS_RETRIES - 1:
                    error(f"  Fallo {url} p{page}: {e}")
                    return all_items
                time.sleep(1.5 ** attempt)

        items = r.json()
        if not items:
            break
        all_items.extend(items)
        bar.update(len(items))

        total_pages = int(r.headers.get("X-WP-TotalPages", 1))
        if page >= total_pages:
            break
        page += 1
        time.sleep(REQUESTS_DELAY)

    bar.close()
    return all_items


def wp_get(endpoint, params=None):
    """GET a la WP REST API (posts, páginas, menús, etc.)."""
    url = f"{SITE_URL}/wp-json/wp/v2/{endpoint}"
    all_items = []
    page = 1
    while True:
        p = {"per_page": 100, "page": page, **(params or {})}
        try:
            r = requests.get(url, params=p, timeout=REQUESTS_TIMEOUT)
            if r.status_code == 400:
                break
            r.raise_for_status()
        except Exception as e:
            warn(f"  wp/{endpoint} p{page}: {e}")
            break
        items = r.json()
        if not items:
            break
        all_items.extend(items)
        total_pages = int(r.headers.get("X-WP-TotalPages", 1))
        if page >= total_pages:
            break
        page += 1
        time.sleep(REQUESTS_DELAY)
    return all_items


def html_get(url):
    headers = {"User-Agent": "Mozilla/5.0 (compatible; SiteExtractor/1.0)"}
    try:
        r = requests.get(url, headers=headers, timeout=REQUESTS_TIMEOUT)
        r.raise_for_status()
        return BeautifulSoup(r.text, "lxml")
    except Exception as e:
        warn(f"  HTML get {url}: {e}")
        return None


def style_header_row(ws, hex_color="1F4E79"):
    fill = PatternFill("solid", fgColor=hex_color)
    font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border


def auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


# ══════════════════════════════════════════════════════
#  PASO 1 – RECONOCIMIENTO HTML
# ══════════════════════════════════════════════════════

def recon_html():
    log("PASO 1 – RECONOCIMIENTO HTML", Fore.MAGENTA)
    soup = html_get(SITE_URL)
    result = {"technology": {}, "menus": [], "banners": [], "colors": [], "fonts": []}

    if not soup:
        warn("No se pudo cargar la página principal vía HTML.")
        return result

    # --- tecnología ---
    gen = soup.find("meta", attrs={"name": "generator"})
    result["technology"]["generator"] = gen["content"] if gen else "Desconocido"

    # WooCommerce
    for script in soup.find_all("script", src=True):
        src = script["src"]
        if "woocommerce" in src.lower():
            result["technology"]["woocommerce"] = True
        if "wp-content" in src.lower():
            result["technology"]["wordpress"] = True

    # --- menús ---
    for nav in soup.find_all("nav"):
        for a in nav.find_all("a", href=True):
            item = {"text": a.get_text(strip=True), "url": a["href"]}
            if item not in result["menus"] and item["text"]:
                result["menus"].append(item)

    # Menú en ul.menu
    for ul in soup.select("ul.menu, ul.nav-menu, #primary-menu li, .main-menu li"):
        for a in ul.find_all("a", href=True):
            item = {"text": a.get_text(strip=True), "url": a["href"]}
            if item not in result["menus"] and item["text"]:
                result["menus"].append(item)

    # --- banners / sliders ---
    for banner in soup.select(".banner, .slider, .hero, [class*='slide'], [class*='banner']"):
        text = banner.get_text(separator=" ", strip=True)
        imgs = [i["src"] for i in banner.find_all("img", src=True)]
        if text or imgs:
            result["banners"].append({"text": text[:300], "images": imgs})

    # --- colores (CSS inline / style blocks) ---
    hex_re = re.compile(r"#(?:[0-9a-fA-F]{3}){1,2}\b")
    style_blocks = [s.string for s in soup.find_all("style") if s.string]
    raw_css = " ".join(style_blocks)
    colors_found = list(set(hex_re.findall(raw_css)))
    result["colors"] = colors_found[:30]

    # --- fuentes ---
    for link in soup.find_all("link", rel=True):
        if "fonts.googleapis.com" in str(link.get("href", "")):
            result["fonts"].append(link["href"])
    for style in soup.find_all("style"):
        if style.string:
            fonts = re.findall(r"font-family:\s*['\"]?([^;'\"]+)", style.string)
            result["fonts"].extend(fonts[:10])
    result["fonts"] = list(set(result["fonts"]))[:15]

    log(f"  Tecnología: {result['technology']}")
    log(f"  Menús encontrados: {len(result['menus'])}")
    log(f"  Banners/sliders: {len(result['banners'])}")
    log(f"  Colores CSS: {len(result['colors'])}")
    return result


# ══════════════════════════════════════════════════════
#  PASO 2 – EXTRACCIÓN WooCommerce API
# ══════════════════════════════════════════════════════

def extract_categories():
    log("Extrayendo categorías…")
    raw = wc_get("products/categories", {"orderby": "id", "hide_empty": False})
    cats = []
    for c in raw:
        cats.append({
            "id":          c.get("id"),
            "name":        c.get("name", ""),
            "slug":        c.get("slug", ""),
            "parent":      c.get("parent", 0),
            "description": BeautifulSoup(c.get("description", ""), "lxml").get_text(),
            "count":       c.get("count", 0),
            "url":         c.get("link", ""),
            "image_url":   (c.get("image") or {}).get("src", ""),
        })
    log(f"  {len(cats)} categorías")
    return cats


def extract_products():
    log("Extrayendo productos…")
    raw = wc_get("products", {"orderby": "id", "status": "publish"})
    products, images = [], []

    for p in raw:
        cats = ", ".join(c["name"] for c in p.get("categories", []))
        tags = ", ".join(t["name"] for t in p.get("tags", []))
        main_img = ""
        if p.get("images"):
            main_img = p["images"][0].get("src", "")
            for img in p["images"]:
                images.append({
                    "product_id":   p["id"],
                    "product_name": p.get("name", ""),
                    "image_id":     img.get("id", ""),
                    "url":          img.get("src", ""),
                    "alt":          img.get("alt", ""),
                    "position":     img.get("position", 0),
                })

        attrs = "; ".join(
            f"{a['name']}: {', '.join(str(o) for o in a.get('options', []))}"
            for a in p.get("attributes", [])
        )

        products.append({
            "id":               p.get("id"),
            "name":             p.get("name", ""),
            "sku":              p.get("sku", ""),
            "status":           p.get("status", ""),
            "type":             p.get("type", ""),
            "price":            p.get("price", ""),
            "regular_price":    p.get("regular_price", ""),
            "sale_price":       p.get("sale_price", ""),
            "stock_status":     p.get("stock_status", ""),
            "stock_quantity":   p.get("stock_quantity"),
            "manage_stock":     p.get("manage_stock", False),
            "categories":       cats,
            "tags":             tags,
            "short_description":BeautifulSoup(p.get("short_description", ""), "lxml").get_text(),
            "description":      BeautifulSoup(p.get("description", ""), "lxml").get_text()[:2000],
            "main_image_url":   main_img,
            "total_images":     len(p.get("images", [])),
            "attributes":       attrs,
            "weight":           p.get("weight", ""),
            "dimensions":       str(p.get("dimensions", {})),
            "url":              p.get("permalink", ""),
            "date_created":     p.get("date_created", ""),
            "date_modified":    p.get("date_modified", ""),
        })

    log(f"  {len(products)} productos, {len(images)} imágenes")
    return products, images


def extract_variations(products_raw_ids):
    log("Extrayendo variantes de productos variables…")
    all_vars = []
    variable_ids = products_raw_ids  # lista de IDs
    for pid in tqdm(variable_ids, desc="  variantes"):
        vars_ = wc_get(f"products/{pid}/variations")
        for v in vars_:
            all_vars.append({
                "product_id":    pid,
                "variation_id":  v.get("id"),
                "sku":           v.get("sku", ""),
                "price":         v.get("price", ""),
                "regular_price": v.get("regular_price", ""),
                "sale_price":    v.get("sale_price", ""),
                "stock_status":  v.get("stock_status", ""),
                "stock_quantity":v.get("stock_quantity"),
                "attributes":    "; ".join(
                    f"{a['name']}: {a.get('option','')}"
                    for a in v.get("attributes", [])
                ),
                "image_url":     (v.get("image") or {}).get("src", ""),
            })
        time.sleep(REQUESTS_DELAY)
    log(f"  {len(all_vars)} variantes")
    return all_vars


def extract_wp_menus():
    log("Extrayendo menús WordPress…")
    menus = wp_get("menus") if False else []  # requiere plugin REST Menus
    # fallback: extraer del HTML
    soup = html_get(SITE_URL)
    menus_html = []
    if soup:
        for a in soup.select("nav a, .menu a, #menu a, .navbar a"):
            text = a.get_text(strip=True)
            href = a.get("href", "")
            parent_li = a.find_parent("li")
            parent_class = " ".join(parent_li.get("class", [])) if parent_li else ""
            if text:
                menus_html.append({
                    "text":    text,
                    "url":     href,
                    "classes": parent_class,
                    "depth":   parent_class.count("sub") + parent_class.count("child"),
                })
    log(f"  {len(menus_html)} ítems de menú")
    return menus_html


def extract_pages_texts():
    log("Extrayendo páginas y textos…")
    pages = wp_get("pages", {"status": "publish"})
    result = []
    for page in pages:
        content = BeautifulSoup(page.get("content", {}).get("rendered", ""), "lxml")
        result.append({
            "id":       page.get("id"),
            "title":    page.get("title", {}).get("rendered", ""),
            "slug":     page.get("slug", ""),
            "url":      page.get("link", ""),
            "excerpt":  BeautifulSoup(page.get("excerpt", {}).get("rendered", ""), "lxml").get_text()[:500],
            "content":  content.get_text(separator=" ", strip=True)[:3000],
            "status":   page.get("status", ""),
            "template": page.get("template", ""),
            "date":     page.get("date", ""),
        })
    log(f"  {len(result)} páginas")
    return result


# ══════════════════════════════════════════════════════
#  PASO 3 – GENERAR EXCEL
# ══════════════════════════════════════════════════════

def generate_excel(products, categories, images, menus, pages, variations, recon):
    log("PASO 3 – Generando Excel…", Fore.MAGENTA)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── HOJA 1: Productos ────────────────────────────
    ws1 = wb.create_sheet("Productos")
    headers1 = [
        "ID", "Nombre", "SKU", "Tipo", "Estado",
        "Precio", "Precio Regular", "Precio Oferta",
        "Stock Status", "Cantidad Stock", "Gestionar Stock",
        "Categorías", "Tags", "Descripción Corta", "Descripción",
        "URL Imagen Principal", "Total Imágenes", "Atributos",
        "Peso", "Dimensiones", "URL Producto",
        "Fecha Creación", "Fecha Modificación"
    ]
    ws1.append(headers1)
    for p in products:
        ws1.append([
            p["id"], p["name"], p["sku"], p["type"], p["status"],
            p["price"], p["regular_price"], p["sale_price"],
            p["stock_status"], p["stock_quantity"], p["manage_stock"],
            p["categories"], p["tags"], p["short_description"], p["description"],
            p["main_image_url"], p["total_images"], p["attributes"],
            p["weight"], p["dimensions"], p["url"],
            p["date_created"], p["date_modified"]
        ])
    style_header_row(ws1)
    auto_width(ws1)
    ws1.freeze_panes = "A2"

    # ── HOJA 2: Variantes ────────────────────────────
    if variations:
        ws_var = wb.create_sheet("Variantes")
        h_var = ["Producto ID", "Variante ID", "SKU", "Precio", "Precio Regular",
                 "Precio Oferta", "Stock Status", "Cantidad Stock", "Atributos", "URL Imagen"]
        ws_var.append(h_var)
        for v in variations:
            ws_var.append([
                v["product_id"], v["variation_id"], v["sku"],
                v["price"], v["regular_price"], v["sale_price"],
                v["stock_status"], v["stock_quantity"],
                v["attributes"], v["image_url"]
            ])
        style_header_row(ws_var)
        auto_width(ws_var)
        ws_var.freeze_panes = "A2"

    # ── HOJA 3: Categorías ───────────────────────────
    ws2 = wb.create_sheet("Categorias")
    headers2 = ["ID", "Nombre", "Slug", "Padre ID", "Descripción", "Conteo", "URL", "URL Imagen"]
    ws2.append(headers2)
    for c in categories:
        ws2.append([
            c["id"], c["name"], c["slug"], c["parent"],
            c["description"], c["count"], c["url"], c["image_url"]
        ])
    style_header_row(ws2)
    auto_width(ws2)
    ws2.freeze_panes = "A2"

    # ── HOJA 4: Imágenes ─────────────────────────────
    ws3 = wb.create_sheet("Imagenes")
    headers3 = ["Producto ID", "Nombre Producto", "Imagen ID", "URL", "Alt Text", "Posición"]
    ws3.append(headers3)
    for img in images:
        ws3.append([
            img["product_id"], img["product_name"],
            img["image_id"], img["url"],
            img["alt"], img["position"]
        ])
    style_header_row(ws3)
    auto_width(ws3)
    ws3.freeze_panes = "A2"

    # ── HOJA 5: Menús y Navegación ───────────────────
    ws4 = wb.create_sheet("Menus")
    headers4 = ["Texto", "URL", "Clases CSS", "Nivel"]
    ws4.append(headers4)
    for m in menus:
        ws4.append([m.get("text",""), m.get("url",""), m.get("classes",""), m.get("depth",0)])
    style_header_row(ws4)
    auto_width(ws4)
    ws4.freeze_panes = "A2"

    # ── HOJA 6: Textos y Banners ─────────────────────
    ws5 = wb.create_sheet("Textos_Banners")
    headers5 = ["Tipo", "ID", "Título", "Slug", "URL", "Resumen", "Contenido", "Estado", "Fecha"]
    ws5.append(headers5)
    for pg in pages:
        ws5.append([
            "Página", pg["id"], pg["title"], pg["slug"], pg["url"],
            pg["excerpt"], pg["content"], pg["status"], pg["date"]
        ])
    for b in recon.get("banners", []):
        ws5.append([
            "Banner", "", "", "", "",
            b.get("text", ""), str(b.get("images", [])), "", ""
        ])
    style_header_row(ws5)
    auto_width(ws5)
    ws5.freeze_panes = "A2"

    # ── HOJA 7: Assets & Tecnología ──────────────────
    ws6 = wb.create_sheet("Assets_Tech")
    ws6.append(["Tipo", "Valor"])
    ws6["A1"].font = Font(bold=True)
    ws6["B1"].font = Font(bold=True)
    tech = recon.get("technology", {})
    for k, v in tech.items():
        ws6.append([f"Tecnología: {k}", str(v)])
    for color in recon.get("colors", []):
        ws6.append(["Color HEX", color])
    for font in recon.get("fonts", []):
        ws6.append(["Fuente", font])
    auto_width(ws6)

    out = os.path.join(OUTPUT_DIR, f"products_{TIMESTAMP}.xlsx")
    wb.save(out)
    log(f"  Excel guardado: {out}")
    return out


# ══════════════════════════════════════════════════════
#  PASO 4 – GENERAR SQL
# ══════════════════════════════════════════════════════

def esc(v):
    if v is None:
        return "NULL"
    s = str(v).replace("\\", "\\\\").replace("'", "\\'").replace("\n", "\\n").replace("\r", "")
    return f"'{s}'"


def generate_sql(products, categories, images, menus, pages, variations):
    log("PASO 4 – Generando SQL…", Fore.MAGENTA)

    lines = [
        "-- ============================================================",
        f"--  BASE DE DATOS: productosdelimpiezalima.com",
        f"--  Generado: {datetime.now().isoformat()}",
        "-- ============================================================",
        "SET NAMES utf8mb4;",
        "SET FOREIGN_KEY_CHECKS=0;",
        "",

        # categorías
        "DROP TABLE IF EXISTS `categories`;",
        """CREATE TABLE `categories` (
  `id`          INT UNSIGNED NOT NULL,
  `name`        VARCHAR(255) NOT NULL,
  `slug`        VARCHAR(255),
  `parent_id`   INT UNSIGNED DEFAULT 0,
  `description` TEXT,
  `count`       INT DEFAULT 0,
  `url`         VARCHAR(512),
  `image_url`   VARCHAR(512),
  PRIMARY KEY (`id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;""",
        "",
    ]

    if categories:
        lines.append("INSERT INTO `categories` (id,name,slug,parent_id,description,count,url,image_url) VALUES")
        rows = []
        for c in categories:
            rows.append(f"  ({c['id']},{esc(c['name'])},{esc(c['slug'])},{c['parent'] or 0},"
                        f"{esc(c['description'])},{c['count'] or 0},{esc(c['url'])},{esc(c['image_url'])})")
        lines.append(",\n".join(rows) + ";")
        lines.append("")

    # productos
    lines += [
        "DROP TABLE IF EXISTS `products`;",
        """CREATE TABLE `products` (
  `id`                INT UNSIGNED NOT NULL,
  `name`              VARCHAR(512) NOT NULL,
  `sku`               VARCHAR(255),
  `type`              VARCHAR(50),
  `status`            VARCHAR(50),
  `price`             DECIMAL(10,2),
  `regular_price`     DECIMAL(10,2),
  `sale_price`        DECIMAL(10,2),
  `stock_status`      VARCHAR(50),
  `stock_quantity`    INT,
  `manage_stock`      TINYINT(1) DEFAULT 0,
  `categories`        TEXT,
  `tags`              TEXT,
  `short_description` TEXT,
  `description`       LONGTEXT,
  `main_image_url`    VARCHAR(1024),
  `total_images`      INT DEFAULT 0,
  `attributes`        TEXT,
  `weight`            VARCHAR(50),
  `dimensions`        VARCHAR(255),
  `url`               VARCHAR(1024),
  `date_created`      DATETIME,
  `date_modified`     DATETIME,
  PRIMARY KEY (`id`),
  FULLTEXT KEY `ft_name_desc` (`name`,`description`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;""",
        "",
    ]

    def to_dec(v):
        try:
            return f"{float(v):.2f}"
        except Exception:
            return "NULL"

    def to_int(v):
        try:
            return int(v)
        except Exception:
            return "NULL"

    def to_dt(v):
        return esc(v[:19].replace("T", " ")) if v and len(v) >= 10 else "NULL"

    if products:
        lines.append("INSERT INTO `products` ("
                     "id,name,sku,type,status,price,regular_price,sale_price,"
                     "stock_status,stock_quantity,manage_stock,categories,tags,"
                     "short_description,description,main_image_url,total_images,"
                     "attributes,weight,dimensions,url,date_created,date_modified) VALUES")
        rows = []
        for p in products:
            rows.append(
                f"  ({p['id']},{esc(p['name'])},{esc(p['sku'])},{esc(p['type'])},{esc(p['status'])},"
                f"{to_dec(p['price'])},{to_dec(p['regular_price'])},{to_dec(p['sale_price'])},"
                f"{esc(p['stock_status'])},{to_int(p['stock_quantity'])},{1 if p['manage_stock'] else 0},"
                f"{esc(p['categories'])},{esc(p['tags'])},"
                f"{esc(p['short_description'])},{esc(p['description'][:3000])},"
                f"{esc(p['main_image_url'])},{p['total_images'] or 0},"
                f"{esc(p['attributes'])},{esc(p['weight'])},{esc(p['dimensions'])},"
                f"{esc(p['url'])},{to_dt(p['date_created'])},{to_dt(p['date_modified'])})"
            )
        lines.append(",\n".join(rows) + ";")
        lines.append("")

    # imágenes
    lines += [
        "DROP TABLE IF EXISTS `product_images`;",
        """CREATE TABLE `product_images` (
  `id`            INT UNSIGNED NOT NULL AUTO_INCREMENT,
  `product_id`    INT UNSIGNED NOT NULL,
  `product_name`  VARCHAR(512),
  `image_id`      INT UNSIGNED,
  `url`           VARCHAR(1024),
  `alt`           VARCHAR(512),
  `position`      INT DEFAULT 0,
  PRIMARY KEY (`id`),
  KEY `idx_product` (`product_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;""",
        "",
    ]

    if images:
        lines.append("INSERT INTO `product_images` (product_id,product_name,image_id,url,alt,position) VALUES")
        rows = []
        for img in images:
            rows.append(
                f"  ({img['product_id']},{esc(img['product_name'])},"
                f"{to_int(img['image_id'])},{esc(img['url'])},{esc(img['alt'])},{img['position'] or 0})"
            )
        lines.append(",\n".join(rows) + ";")
        lines.append("")

    # variantes
    if variations:
        lines += [
            "DROP TABLE IF EXISTS `product_variations`;",
            """CREATE TABLE `product_variations` (
  `id`             INT UNSIGNED NOT NULL AUTO_INCREMENT,
  `product_id`     INT UNSIGNED NOT NULL,
  `variation_id`   INT UNSIGNED,
  `sku`            VARCHAR(255),
  `price`          DECIMAL(10,2),
  `regular_price`  DECIMAL(10,2),
  `sale_price`     DECIMAL(10,2),
  `stock_status`   VARCHAR(50),
  `stock_quantity` INT,
  `attributes`     TEXT,
  `image_url`      VARCHAR(1024),
  PRIMARY KEY (`id`),
  KEY `idx_product` (`product_id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;""",
            "",
        ]
        lines.append("INSERT INTO `product_variations` ("
                     "product_id,variation_id,sku,price,regular_price,sale_price,"
                     "stock_status,stock_quantity,attributes,image_url) VALUES")
        rows = []
        for v in variations:
            rows.append(
                f"  ({v['product_id']},{to_int(v['variation_id'])},{esc(v['sku'])},"
                f"{to_dec(v['price'])},{to_dec(v['regular_price'])},{to_dec(v['sale_price'])},"
                f"{esc(v['stock_status'])},{to_int(v['stock_quantity'])},"
                f"{esc(v['attributes'])},{esc(v['image_url'])})"
            )
        lines.append(",\n".join(rows) + ";")
        lines.append("")

    # menús
    lines += [
        "DROP TABLE IF EXISTS `nav_menus`;",
        """CREATE TABLE `nav_menus` (
  `id`      INT UNSIGNED NOT NULL AUTO_INCREMENT,
  `text`    VARCHAR(512),
  `url`     VARCHAR(1024),
  `classes` VARCHAR(512),
  `depth`   INT DEFAULT 0,
  PRIMARY KEY (`id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;""",
        "",
    ]
    if menus:
        lines.append("INSERT INTO `nav_menus` (text,url,classes,depth) VALUES")
        rows = [f"  ({esc(m.get('text',''))},{esc(m.get('url',''))},{esc(m.get('classes',''))},{m.get('depth',0)})"
                for m in menus]
        lines.append(",\n".join(rows) + ";")
        lines.append("")

    # páginas
    lines += [
        "DROP TABLE IF EXISTS `pages`;",
        """CREATE TABLE `pages` (
  `id`       INT UNSIGNED NOT NULL,
  `title`    VARCHAR(512),
  `slug`     VARCHAR(255),
  `url`      VARCHAR(1024),
  `excerpt`  TEXT,
  `content`  LONGTEXT,
  `status`   VARCHAR(50),
  `template` VARCHAR(255),
  `date`     DATETIME,
  PRIMARY KEY (`id`)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;""",
        "",
    ]
    if pages:
        lines.append("INSERT INTO `pages` (id,title,slug,url,excerpt,content,status,template,date) VALUES")
        rows = [
            f"  ({pg['id']},{esc(pg['title'])},{esc(pg['slug'])},{esc(pg['url'])},"
            f"{esc(pg['excerpt'])},{esc(pg['content'][:3000])},{esc(pg['status'])},"
            f"{esc(pg['template'])},{to_dt(pg['date'])})"
            for pg in pages
        ]
        lines.append(",\n".join(rows) + ";")

    lines += ["", "SET FOREIGN_KEY_CHECKS=1;", "-- FIN"]

    out = os.path.join(OUTPUT_DIR, f"database_{TIMESTAMP}.sql")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    log(f"  SQL guardado: {out}")
    return out


# ══════════════════════════════════════════════════════
#  PASO 5 – GENERAR JSON
# ══════════════════════════════════════════════════════

def generate_json_files(products, categories, images, menus, pages, recon):
    log("PASO 5 – Generando JSON…", Fore.MAGENTA)

    images_out = os.path.join(OUTPUT_DIR, "images.json")
    with open(images_out, "w", encoding="utf-8") as f:
        json.dump(images, f, ensure_ascii=False, indent=2)
    log(f"  images.json  → {len(images)} registros")

    sitemap = {
        "generated": datetime.now().isoformat(),
        "site": SITE_URL,
        "total_products": len(products),
        "total_categories": len(categories),
        "total_images": len(images),
        "total_pages": len(pages),
        "pages": [{"id": pg["id"], "title": pg["title"], "url": pg["url"], "slug": pg["slug"]} for pg in pages],
        "categories": [{"id": c["id"], "name": c["name"], "url": c["url"], "count": c["count"]} for c in categories],
        "menus": menus,
        "technology": recon.get("technology", {}),
        "colors": recon.get("colors", []),
        "fonts": recon.get("fonts", []),
    }
    sitemap_out = os.path.join(OUTPUT_DIR, "sitemap.json")
    with open(sitemap_out, "w", encoding="utf-8") as f:
        json.dump(sitemap, f, ensure_ascii=False, indent=2)
    log(f"  sitemap.json → guardado")

    products_out = os.path.join(OUTPUT_DIR, "products.json")
    with open(products_out, "w", encoding="utf-8") as f:
        json.dump(products, f, ensure_ascii=False, indent=2)
    log(f"  products.json → {len(products)} productos")

    cats_out = os.path.join(OUTPUT_DIR, "categories.json")
    with open(cats_out, "w", encoding="utf-8") as f:
        json.dump(categories, f, ensure_ascii=False, indent=2)
    log(f"  categories.json → {len(categories)} categorías")

    return images_out, sitemap_out


# ══════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════

def main():
    print(f"\n{Fore.MAGENTA}{'='*60}")
    print(f"  EXTRACTOR COMPLETO – productosdelimpiezalima.com")
    print(f"  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"{'='*60}{Style.RESET_ALL}\n")

    # ── PASO 1 ───────────────────────────────────────
    recon = recon_html()
    print()

    # ── PASO 2 ───────────────────────────────────────
    log("PASO 2 – EXTRACCIÓN DE DATOS vía WooCommerce API", Fore.MAGENTA)
    categories = extract_categories()
    products, images = extract_products()

    variable_ids = []
    raw = wc_get("products", {"type": "variable", "per_page": PER_PAGE})
    variable_ids = [p["id"] for p in raw]
    variations = extract_variations(variable_ids) if variable_ids else []

    menus   = extract_wp_menus()
    pages   = extract_pages_texts()
    print()

    # ── PASO 3 ───────────────────────────────────────
    excel_file = generate_excel(products, categories, images, menus, pages, variations, recon)
    print()

    # ── PASO 4 ───────────────────────────────────────
    sql_file = generate_sql(products, categories, images, menus, pages, variations)
    print()

    # ── PASO 5 ───────────────────────────────────────
    generate_json_files(products, categories, images, menus, pages, recon)
    print()

    # ── RESUMEN ──────────────────────────────────────
    print(f"{Fore.GREEN}{'='*60}")
    print(f"  EXTRACCIÓN COMPLETA")
    print(f"  Productos:   {len(products)}")
    print(f"  Variantes:   {len(variations)}")
    print(f"  Categorías:  {len(categories)}")
    print(f"  Imágenes:    {len(images)}")
    print(f"  Páginas:     {len(pages)}")
    print(f"  Menús:       {len(menus)}")
    print(f"  Archivos en: {os.path.abspath(OUTPUT_DIR)}/")
    print(f"{'='*60}{Style.RESET_ALL}\n")


if __name__ == "__main__":
    main()
