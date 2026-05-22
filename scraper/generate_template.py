"""Genera el archivo Excel plantilla con las hojas y encabezados listos."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = os.path.join(os.path.dirname(__file__), "..", "data", "products_template.xlsx")
os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)

SHEETS = {
    "Productos": [
        "ID","Nombre","SKU","Tipo","Estado",
        "Precio","Precio Regular","Precio Oferta",
        "Stock Status","Cantidad Stock","Gestionar Stock",
        "Categorías","Tags","Descripción Corta","Descripción",
        "URL Imagen Principal","Total Imágenes","Atributos",
        "Peso","Dimensiones","URL Producto",
        "Fecha Creación","Fecha Modificación"
    ],
    "Variantes": [
        "Producto ID","Variante ID","SKU","Precio","Precio Regular",
        "Precio Oferta","Stock Status","Cantidad Stock","Atributos","URL Imagen"
    ],
    "Categorias": [
        "ID","Nombre","Slug","Padre ID","Descripción","Conteo","URL","URL Imagen"
    ],
    "Imagenes": [
        "Producto ID","Nombre Producto","Imagen ID","URL","Alt Text","Posición"
    ],
    "Menus": [
        "Texto","URL","Clases CSS","Nivel"
    ],
    "Textos_Banners": [
        "Tipo","ID","Título","Slug","URL","Resumen","Contenido","Estado","Fecha"
    ],
    "Assets_Tech": [
        "Tipo","Valor"
    ],
}

def style_header(ws, hex_color="1F4E79"):
    fill   = PatternFill("solid", fgColor=hex_color)
    font   = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    align  = Alignment(horizontal="center", vertical="center")
    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill   = fill
        cell.font   = font
        cell.alignment = align
        cell.border = border
    ws.row_dimensions[1].height = 20

wb = openpyxl.Workbook()
wb.remove(wb.active)

for sheet_name, headers in SHEETS.items():
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    style_header(ws)
    ws.freeze_panes = "A2"
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    # fila ejemplo (vacía con indicación)
    ws.append(["← Ejecuta woo_extractor.py para poblar esta hoja"] + [""] * (len(headers)-1))
    ws["A2"].font = Font(italic=True, color="888888")

wb.save(OUTPUT)
print(f"Plantilla guardada: {os.path.abspath(OUTPUT)}")
